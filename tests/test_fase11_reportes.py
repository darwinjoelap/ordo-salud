"""Fase 11: reportes (cifras, alcance por rol, PDF y Excel) y perfil del médico estilo Pediae."""
from datetime import time
from decimal import Decimal
from io import BytesIO

from django.utils import timezone

from apps.agenda.models import Cita
from apps.clinica.consultas import abrir_consulta
from apps.clinica.models import DiagnosticoConsulta
from apps.cobros import servicios as cobros
from apps.core.permisos import Rol
from apps.equipo import servicios as equipo
from apps.equipo.models import Membresia
from apps.pacientes.models import Paciente
from apps.plataforma.models import Especialidad

from .base import BaseTest


class ReportesTest(BaseTest):
    def setUp(self):
        super().setUp()
        esp = Especialidad.objects.get(codigo='medicina_general')
        actor = Membresia.objects.get(usuario=self.titular)
        m, clave = equipo.agregar(self.tenant, actor, 'drdos', 'Luis', 'Mar', Rol.MEDICO, especialidad=esp)
        m.usuario.debe_cambiar_clave = False
        m.usuario.save()
        self.medico = self.entrar(self.tenant, 'drdos', clave)
        hoy = timezone.localdate()
        with self.en_tenant():
            self.p_titular = actor.perfil
            self.p_medico = m.perfil
            ana = Paciente.objects.create(nombres='Ana', apellidos='Uno', sexo='F', fecha_nacimiento='2020-01-01')
            beto = Paciente.objects.create(nombres='Beto', apellidos='Dos', sexo='M', fecha_nacimiento='1980-05-05')
            c1 = abrir_consulta(paciente=ana, profesional=self.p_titular)
            DiagnosticoConsulta.objects.create(consulta=c1, paciente=ana, codigo='J00', descripcion='Resfriado común')
            c2 = abrir_consulta(paciente=beto, profesional=self.p_titular, tipo='CONTROL')
            DiagnosticoConsulta.objects.create(consulta=c2, paciente=beto, codigo='J00', descripcion='Resfriado común')
            abrir_consulta(paciente=beto, profesional=self.p_medico)
            Cita.objects.create(paciente=ana, profesional=self.p_titular, especialidad=self.p_titular.especialidad,
                                fecha=hoy, hora_inicio=time(8), hora_fin=time(8, 30), estado='ATENDIDA', consulta=c1)
            Cita.objects.create(paciente=beto, profesional=self.p_titular, especialidad=self.p_titular.especialidad,
                                fecha=hoy, hora_inicio=time(9), hora_fin=time(9, 30), estado='NO_ASISTIO')
            cargo = cobros.crear_cargo(paciente=ana, descripcion='Consulta', precio_usd=Decimal('40'), consulta=c1)
            cobros.registrar_pago(paciente=ana, monto=Decimal('40'), moneda='USD', metodo='ZELLE')
            self.cargo = cargo

    def test_cifras_del_consultorio(self):
        r = self.cliente.get('/reportes/')
        self.assertEqual(r.status_code, 200)
        c = r.context
        self.assertEqual(c['consultas'], 3)
        self.assertEqual(c['pacientes_atendidos'], 2)
        self.assertEqual(c['pacientes_nuevos'], 2)
        self.assertEqual(c['agenda']['tasa_asistencia'], 50.0)
        self.assertEqual(c['agenda']['sin_cita'], 2)
        self.assertEqual(c['morbilidad'][0]['valor'], 2)
        self.assertEqual(c['finanzas']['cobrado'], Decimal('40.00'))
        self.assertEqual(len(c['por_profesional']), 2)
        self.assertContains(r, 'Resfriado común')

    def test_filtro_por_profesional(self):
        r = self.cliente.get(f'/reportes/?profesional={self.p_medico.pk}')
        self.assertEqual(r.context['consultas'], 1)
        self.assertEqual(r.context['finanzas']['cobrado'], Decimal('0'))
        self.assertEqual(r.context['por_profesional'], [])

    def test_medico_ve_solo_lo_suyo(self):
        r = self.medico.get(f'/reportes/?profesional={self.p_titular.pk}')     # el filtro ajeno se ignora
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context['consultas'], 1)
        self.assertContains(r, 'Mi gestión')
        self.assertNotContains(r, 'Por profesional')

    def test_asistente_sin_reportes(self):
        actor = Membresia.objects.get(usuario=self.titular)
        m, clave = equipo.agregar(self.tenant, actor, 'caja', 'Ana', 'Gil', Rol.ASISTENTE)
        m.usuario.debe_cambiar_clave = False
        m.usuario.save()
        self.assertEqual(self.entrar(self.tenant, 'caja', clave).get('/reportes/').status_code, 403)

    def test_exporta_pdf_y_excel(self):
        r = self.cliente.get('/reportes/pdf/')
        self.assertEqual(r['Content-Type'], 'application/pdf')
        from pypdf import PdfReader
        texto = ' '.join(p.extract_text() for p in PdfReader(BytesIO(r.content)).pages)
        self.assertIn('REPORTE DE GESTIÓN', texto)
        self.assertIn('Resfriado común', texto)
        r = self.cliente.get('/reportes/excel/')
        self.assertIn('spreadsheetml', r['Content-Type'])
        from openpyxl import load_workbook
        libro = load_workbook(BytesIO(r.content))
        self.assertEqual(libro.sheetnames[:3], ['Resumen', 'Por mes', 'Enfermedades'])
        self.assertEqual(libro['Consultas'].max_row, 4 + 3)

    def test_documento_de_muestra_del_medico(self):
        r = self.medico.get('/documentos/muestra/')
        self.assertEqual(r['Content-Type'], 'application/pdf')
        # un médico no ve la muestra de otro; el titular (gestiona el equipo) sí
        self.assertEqual(self.medico.get(f'/documentos/muestra/{self.p_titular.pk}/').status_code, 403)
        self.assertEqual(self.cliente.get(f'/documentos/muestra/{self.p_medico.pk}/').status_code, 200)

    def test_perfil_sin_marca_de_agua(self):
        from apps.core.pdf import marca_agua_de
        with self.en_tenant():
            self.p_medico.marca_agua = False
            self.p_medico.save()
            self.assertIsNone(marca_agua_de(self.tenant, self.p_medico))
