"""
Reporte de ginecología y obstetricia (estadísticas de Ginea `reportes/views.py::estadisticas` + mejoras): tipos de
consulta, controles prenatales y EG promedio, embarazos en curso y partos esperados, grupos de edad, métodos
anticonceptivos, VIH, VPH, citología al día/vencida, antecedentes familiares y citas por día. Mismos filtros y alcance
que «Reportes» (la médica ve lo suyo; titular/administración, todo el consultorio). Los datos sensibles solo se
cuentan: el reporte no muestra nombres salvo en la lista de partos esperados (agenda clínica de la médica).
"""
from collections import Counter
from datetime import timedelta
from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone

from apps.core import auditoria
from apps.reportes import selectors as rs

from . import obstetricia as ob
from .models import ControlPrenatal
from .paneles import CITOLOGIA_VIGENTE_DIAS, citologia_alterada
from .secciones import FAMILIARES, TIPOS_CONSULTA

GRUPOS_EDAD = [('Hasta 20 años', 0, 21), ('21 a 30', 21, 31), ('31 a 40', 31, 41), ('41 a 50', 41, 51),
               ('51 a 60', 51, 61), ('Más de 60', 61, 200)]
DIAS = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado', 'Domingo']


def datos_reporte(f):
    from apps.agenda.models import Cita
    from apps.clinica.models import RegistroSeccion, RegistroSeccionPaciente
    hoy = timezone.localdate()
    consultas = rs.consultas_de(f).select_related('paciente')
    ids = list(consultas.values_list('pk', flat=True))
    pacientes = {c.paciente_id: c.paciente for c in consultas}

    tipos = Counter(r.datos.get('tipo') for r in RegistroSeccion.objects.filter(
        consulta_id__in=ids, seccion__codigo='go.tipo'))
    controles = list(ControlPrenatal.objects.filter(consulta_id__in=ids))
    eg = [c.eg_dias for c in controles if c.eg_dias is not None]

    edades = Counter()
    for p in pacientes.values():
        if p.edad_anios is not None:
            edades[next(e for e, a, b in GRUPOS_EDAD if a <= p.edad_anios < b)] += 1

    fichas = RegistroSeccionPaciente.objects.filter(paciente_id__in=pacientes.keys(), vigente=True)
    metodos, familiares, vih = Counter(), Counter(), Counter()
    vph_vacuna = vph_dx = cito_al_dia = cito_vencida = cito_sin = 0
    con_cito = set()
    for r in fichas.filter(seccion__codigo='go.planificacion'):
        metodo = (r.datos.get('metodo') or '').strip()
        if metodo:
            metodos[metodo[:1].upper() + metodo[1:].lower()] += 1     # texto libre (como en Ginea)
    for r in fichas.filter(seccion__codigo='go.familiares'):
        familiares.update(r.datos.get('antecedentes') or [])
    vph_no = its = cito_alterada = 0
    condiciones = Counter()
    for r in fichas.filter(seccion__codigo='go.gineco'):
        d = r.datos
        vph_vacuna += d.get('vph_vacuna') is True
        vph_no += d.get('vph_vacuna') is False
        vph_dx += bool(d.get('vph_diagnostico'))
        its += bool((d.get('its_previas') or '').strip())
        cito_alterada += citologia_alterada(d.get('citologia_resultado'))
        for clave, nombre in (('menopausia', 'Menopausia'), ('histerectomia', 'Histerectomía'),
                              ('dispareunia', 'Dispareunia'), ('ciclo_regular', None)):
            if clave == 'ciclo_regular':
                if d.get(clave) is False:
                    condiciones['Ciclo irregular'] += 1
            elif d.get(clave):
                condiciones[nombre] += 1
        if d.get('vih_resultado'):
            vih[d['vih_resultado']] += 1
        fecha = ob.a_fecha(d.get('citologia_fecha'))
        if fecha:
            con_cito.add(r.paciente_id)
            if (hoy - fecha).days <= CITOLOGIA_VIGENTE_DIAS:
                cito_al_dia += 1
            else:
                cito_vencida += 1
    cito_sin = len([p for p in pacientes if p not in con_cito])
    habitos = Counter()
    for r in fichas.filter(seccion__codigo='go.habitos'):
        for clave, nombre in (('tabaquismo', 'Tabaquismo'), ('alcohol', 'Alcoholismo'),
                              ('transfusiones', 'Transfusiones')):
            habitos[nombre] += bool(r.datos.get(clave))
    for r in fichas.filter(seccion__codigo='go.planificacion'):
        if r.datos.get('deseo_embarazo'):
            condiciones['Desea embarazo'] += 1
    con_ficha = fichas.filter(seccion__codigo='go.gineco').count()

    # Embarazos en curso (de las pacientes atendidas) y partos esperados en los próximos 30 días
    ultimos = {}
    for c in ControlPrenatal.objects.filter(paciente_id__in=pacientes.keys()).order_by('fecha', 'pk'):
        ultimos[c.paciente_id] = c
    en_curso = [c for c in ultimos.values() if c.en_curso]
    proximos = sorted((c for c in en_curso if c.fpp and hoy <= c.fpp <= hoy + timedelta(days=30)),
                      key=lambda c: c.fpp)
    trimestres = Counter(ob.trimestre(c.eg_dias_en(hoy)) for c in en_curso)

    citas = rs._aplicar(Cita.objects.filter(fecha__range=(f.desde, f.hasta)), f)
    por_dia = Counter(c.fecha.weekday() for c in citas.only('fecha'))
    nombres_fam = dict(next(c for c in FAMILIARES.esquema['campos'] if c['clave'] == 'antecedentes')['opciones'])
    return {
        'consultas': len(ids), 'pacientes': len(pacientes), 'prenatales': len(controles),
        'eg_promedio': ob.texto_eg(round(sum(eg) / len(eg))) if eg else '',
        'en_curso': len(en_curso), 'proximos': proximos,
        'por_tipo': rs._barras([(n, tipos.get(k, 0)) for k, n in TIPOS_CONSULTA if tipos.get(k)]),
        'trimestres': rs._barras([(f'{t}.er trimestre' if t in (1, 3) else f'{t}.º trimestre', trimestres.get(t, 0))
                                  for t in (1, 2, 3)]),
        'edades': rs._barras([(e, edades.get(e, 0)) for e, _, _ in GRUPOS_EDAD]),
        'metodos': rs._barras(metodos.most_common(8)),
        'familiares': rs._barras([(nombres_fam.get(k, k), v) for k, v in familiares.most_common()]),
        'vih': rs._barras([(n, vih.get(k, 0)) for k, n in (('NEGATIVO', 'Negativo'), ('POSITIVO', 'Positivo'),
                                                          ('NO_REALIZADO', 'No realizada')) if vih.get(k)]),
        'citologia': rs._barras([('Al día (último año)', cito_al_dia), ('Vencida', cito_vencida),
                                 ('Sin registro', cito_sin)]),
        'vph': {'vacunadas': vph_vacuna, 'diagnostico': vph_dx},
        'vph_barras': rs._barras([('Vacunadas', vph_vacuna), ('No vacunadas', vph_no),
                                  ('Sin dato', max(con_ficha - vph_vacuna - vph_no, 0))]),
        'infecciones': rs._barras([('VIH positivo', vih.get('POSITIVO', 0)), ('Diagnóstico de VPH', vph_dx),
                                   ('Con ITS previas', its), ('Citología alterada', cito_alterada)]),
        'condiciones': rs._barras(condiciones.most_common()),
        'habitos': rs._barras([(k, v) for k, v in habitos.items() if v]),
        'morbilidad': rs.morbilidad_de(consultas),
        'con_ficha': con_ficha,
        'por_dia': rs._barras([(DIAS[i], por_dia.get(i, 0)) for i in range(7)]),
    }


@login_required
def tablero(request):
    if not rs.puede_ver_reportes(request):
        raise PermissionDenied
    f = rs.filtros_de(request)
    datos = datos_reporte(f)
    if request.GET.get('formato') == 'excel':
        auditoria.registrar(request, 'EXPORTAR', 'reportes', detalle=f'Reporte de ginecología {f.periodo_texto}')
        return _excel(datos, f)
    return render(request, 'ginecologia/reporte.html', {
        'titulo': 'Reporte de ginecología', 'filtros': f, **datos, 'rangos': rs.rangos_rapidos(),
        'pestanas': rs.pestanas(request, 'go_reporte'),
        'query': request.GET.urlencode()})


def _excel(datos, f):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Ginecología'
    hoja.append([f'Reporte de ginecología y obstetricia · {f.periodo_texto}'])
    hoja['A1'].font = Font(bold=True, size=13)
    for etiqueta, valor in (('Consultas', datos['consultas']), ('Pacientes atendidas', datos['pacientes']),
                            ('Controles prenatales', datos['prenatales']), ('EG promedio', datos['eg_promedio']),
                            ('Embarazos en curso', datos['en_curso']),
                            ('Vacunadas contra VPH', datos['vph']['vacunadas']),
                            ('Con diagnóstico de VPH', datos['vph']['diagnostico'])):
        hoja.append([etiqueta, valor])
    for titulo, clave in [('Enfermedades más frecuentes', 'morbilidad'), ('Infecciones y hallazgos', 'infecciones'),
                          ('Vacuna contra el VPH', 'vph_barras'), ('Condiciones gineco-obstétricas', 'condiciones'),
                          ('Hábitos', 'habitos'),
                          ('Tipo de consulta', 'por_tipo'), ('Embarazos por trimestre', 'trimestres'),
                          ('Grupos de edad', 'edades'), ('Métodos anticonceptivos', 'metodos'),
                          ('Citología', 'citologia'), ('Prueba de VIH', 'vih'),
                          ('Antecedentes familiares', 'familiares'), ('Citas por día', 'por_dia')]:
        hoja.append([])
        hoja.append([titulo])
        hoja.cell(hoja.max_row, 1).font = Font(bold=True)
        for b in datos[clave]:
            hoja.append([b['etiqueta'], b['valor']])
    if datos['proximos']:
        hoja.append([])
        hoja.append(['Partos esperados (próximos 30 días)'])
        hoja.cell(hoja.max_row, 1).font = Font(bold=True)
        for c in datos['proximos']:
            hoja.append([c.paciente.nombre_completo, c.fpp.strftime('%d/%m/%Y')])
    hoja.column_dimensions['A'].width = 42
    salida = BytesIO()
    libro.save(salida)
    r = HttpResponse(salida.getvalue(),
                     content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    r['Content-Disposition'] = f'attachment; filename="reporte_ginecologia_{f.desde}_{f.hasta}.xlsx"'
    return r
