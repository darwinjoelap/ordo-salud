"""
Reporte pediátrico (estadísticas de Pediae `reportes/views.py::estadisticas` + mejoras): tipos de consulta, estado
nutricional y clasificación OMS, talla baja, grupos de edad pediátricos, antecedentes familiares, cobertura de
vacunas, lactancia materna, pesquisas y citas por día de la semana. Usa los mismos filtros y alcance que «Reportes»
(el médico ve lo suyo; titular/administración, todo el consultorio).
"""
from collections import Counter
from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import render

from apps.core import auditoria
from apps.core.fechas import edad_partes
from apps.reportes import selectors as rs

from . import vacunas
from .models import MedicionCrecimiento, PesquisaRealizada
from .secciones import FAMILIARES, TIPOS_CONSULTA

GRUPOS_EDAD = [('0 a 1 año', 0, 1), ('1 a 3 años', 1, 3), ('3 a 6 años', 3, 6), ('6 a 10 años', 6, 10),
               ('10 a 15 años', 10, 15), ('15 años o más', 15, 200)]
ORDEN_NUTRICIONAL = ['Desnutrición severa (<p3)', 'Desnutrición (p3–p10)', 'Riesgo de desnutrición (p10–p15)',
                     'Eutrófico (p15–p85)', 'Sobrepeso (p85–p97)', 'Obesidad (>p97)']
DIAS = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado', 'Domingo']


def datos_reporte(f):
    from apps.agenda.models import Cita
    from apps.clinica.models import RegistroSeccion, RegistroSeccionPaciente
    consultas = rs.consultas_de(f).select_related('paciente')
    ids = list(consultas.values_list('pk', flat=True))
    pacientes = {c.paciente_id: c.paciente for c in consultas}

    tipos = Counter(r.datos.get('tipo') for r in RegistroSeccion.objects.filter(
        consulta_id__in=ids, seccion__codigo='ped.tipo'))
    etiquetas_tipo = dict(TIPOS_CONSULTA)
    por_tipo = [(etiquetas_tipo[k], tipos.get(k, 0)) for k in etiquetas_tipo if tipos.get(k)]

    # Estado nutricional: la última medición de cada paciente en el período
    ultimas = {}
    for m in MedicionCrecimiento.objects.filter(consulta_id__in=ids).order_by('fecha', 'pk'):
        ultimas[m.paciente_id] = m
    nutricion = Counter(m.estado_nutricional for m in ultimas.values() if m.estado_nutricional)
    imc = Counter(m.clasificacion_imc for m in ultimas.values() if m.clasificacion_imc)
    talla = Counter(m.clasificacion_talla for m in ultimas.values() if m.clasificacion_talla)

    edades = Counter()                  # edad de cada paciente en su primera consulta del período
    vistos = set()
    for c in consultas.order_by('fecha_hora'):
        if c.paciente_id in vistos or not c.paciente.fecha_nacimiento:
            continue
        vistos.add(c.paciente_id)
        anios = (edad_partes(c.paciente.fecha_nacimiento, c.fecha_hora) or (0,))[0]
        etiqueta = next(e for e, a, b in GRUPOS_EDAD if a <= anios < b)
        edades[etiqueta] += 1

    fichas = RegistroSeccionPaciente.objects.filter(paciente_id__in=pacientes.keys(), vigente=True)
    familiares = Counter()
    lme_si = lme_no = 0
    lme_meses = []
    for r in fichas.filter(seccion__codigo='ped.familiares'):
        familiares.update(r.datos.get('antecedentes') or [])
    for r in fichas.filter(seccion__codigo='ped.alimentacion'):
        if r.datos.get('lme') is True:
            lme_si += 1
            if r.datos.get('lme_meses'):
                lme_meses.append(r.datos['lme_meses'])
        elif r.datos.get('lme') is False:
            lme_no += 1
    nombres_fam = dict(next(c for c in FAMILIARES.esquema['campos'] if c['clave'] == 'antecedentes')['opciones'])

    menores = [p for p in pacientes.values() if p.fecha_nacimiento and (p.edad_anios or 0) < 19]
    al_dia = atrasados = 0
    for p in menores:
        r = vacunas.resumen(p)
        if r['atrasadas']:
            atrasados += 1
        elif r['al_dia']:
            al_dia += 1

    pesquisas = Counter(PesquisaRealizada.objects.filter(paciente_id__in=pacientes.keys())
                        .values_list('pesquisa__nombre', flat=True))
    citas = rs._aplicar(Cita.objects.filter(fecha__range=(f.desde, f.hasta)), f)
    por_dia = Counter(c.fecha.weekday() for c in citas.only('fecha'))
    return {
        'consultas': len(ids), 'pacientes': len(pacientes), 'con_medicion': len(ultimas),
        'morbilidad': rs.morbilidad_de(consultas),
        'por_tipo': rs._barras(por_tipo),
        'nutricion': rs._barras([(e, nutricion[e]) for e in ORDEN_NUTRICIONAL if nutricion.get(e)]),
        'imc': rs._barras(imc.most_common()), 'talla': rs._barras(talla.most_common()),
        'edades': rs._barras([(e, edades.get(e, 0)) for e, _, _ in GRUPOS_EDAD]),
        'familiares': rs._barras([(nombres_fam.get(k, k), v) for k, v in familiares.most_common()]),
        'vacunas': {'menores': len(menores), 'al_dia': al_dia, 'atrasados': atrasados,
                    'pct_al_dia': round(al_dia * 100 / len(menores)) if menores else None},
        'lactancia': {'si': lme_si, 'no': lme_no, 'pct': round(lme_si * 100 / (lme_si + lme_no))
                      if lme_si + lme_no else None,
                      'promedio_meses': round(sum(lme_meses) / len(lme_meses), 1) if lme_meses else None},
        'pesquisas': rs._barras(pesquisas.most_common()),
        'por_dia': rs._barras([(DIAS[i], por_dia.get(i, 0)) for i in range(7)]),
    }


@login_required
def tablero(request):
    if not rs.puede_ver_reportes(request):
        raise PermissionDenied
    f = rs.filtros_de(request)
    datos = datos_reporte(f)
    if request.GET.get('formato') == 'excel':
        auditoria.registrar(request, 'EXPORTAR', 'reportes', detalle=f'Reporte pediátrico {f.periodo_texto}')
        return _excel(datos, f)
    query = request.GET.urlencode()
    return render(request, 'pediatria/reporte.html', {
        'titulo': 'Reporte pediátrico', 'filtros': f, **datos, 'rangos': rs.rangos_rapidos(),
        'pestanas': rs.pestanas(request, 'ped_reporte'),
        'query': query})


def _excel(datos, f):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Pediatría'
    hoja.append([f'Reporte pediátrico · {f.periodo_texto}'])
    hoja['A1'].font = Font(bold=True, size=13)
    hoja.append(['Consultas', datos['consultas']])
    hoja.append(['Pacientes atendidos', datos['pacientes']])
    for titulo, clave in [('Enfermedades más frecuentes', 'morbilidad'), ('Tipo de consulta', 'por_tipo'),
                          ('Estado nutricional (peso/edad)', 'nutricion'),
                          ('Clasificación OMS (IMC/edad)', 'imc'), ('Talla para la edad', 'talla'),
                          ('Grupos de edad', 'edades'), ('Antecedentes familiares', 'familiares'),
                          ('Pesquisas realizadas', 'pesquisas'), ('Citas por día', 'por_dia')]:
        hoja.append([])
        hoja.append([titulo])
        hoja.cell(hoja.max_row, 1).font = Font(bold=True)
        for b in datos[clave]:
            hoja.append([b['etiqueta'], b['valor']])
    v, lac = datos['vacunas'], datos['lactancia']
    hoja.append([])
    hoja.append(['Vacunas: menores atendidos', v['menores']])
    hoja.append(['Con esquema al día', v['al_dia']])
    hoja.append(['Con dosis atrasadas', v['atrasados']])
    hoja.append(['Lactancia materna exclusiva: sí / no', lac['si'], lac['no']])
    hoja.column_dimensions['A'].width = 42
    salida = BytesIO()
    libro.save(salida)
    r = HttpResponse(salida.getvalue(),
                     content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    r['Content-Disposition'] = f'attachment; filename="reporte_pediatrico_{f.desde}_{f.hasta}.xlsx"'
    return r
