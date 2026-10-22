"""Exportaciones del reporte: PDF (ReportLab, con el encabezado del consultorio) y Excel (openpyxl)."""
from io import BytesIO

from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Spacer, Table, TableStyle

from apps.core.pdf import ESTILO_CELDA, ESTILOS, Paragraph, documento, encabezado_tenant, pie_tenant
from apps.core.templatetags.ordo import usd

PRIMARIO = colors.HexColor('#145F8F')
ARENA = colors.HexColor('#F8F4EE')
BORDE = colors.HexColor('#DCE4EB')
ANCHO = 186 * mm


def _titulo(texto):
    from reportlab.lib.styles import ParagraphStyle
    estilo = ParagraphStyle('s', parent=ESTILOS['Normal'], fontName='Helvetica-Bold', fontSize=10.5, leading=13,
                            textColor=PRIMARIO, spaceBefore=8, spaceAfter=3)
    return Paragraph(texto, estilo)


def _tabla(encabezado, filas, anchos, alinear_derecha=()):
    datos = [[Paragraph(f'<b>{h}</b>', ESTILO_CELDA) for h in encabezado]]
    datos += [[Paragraph(str(c), ESTILO_CELDA) for c in fila] for fila in filas]
    t = Table(datos, colWidths=anchos, repeatRows=1)
    estilo = [('BACKGROUND', (0, 0), (-1, 0), ARENA), ('LINEBELOW', (0, 0), (-1, -1), 0.4, BORDE),
              ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('TOPPADDING', (0, 0), (-1, -1), 3),
              ('BOTTOMPADDING', (0, 0), (-1, -1), 3)]
    for col in alinear_derecha:
        estilo.append(('ALIGN', (col, 0), (col, -1), 'RIGHT'))
    t.setStyle(TableStyle(estilo))
    return t


def _kpis(pares):
    """Fila de cifras: [(etiqueta, valor)] en cajas."""
    from reportlab.lib.styles import ParagraphStyle
    cifra = ParagraphStyle('cifra', parent=ESTILO_CELDA, fontName='Helvetica-Bold', fontSize=15, leading=19,
                           textColor=colors.HexColor('#0C385B'))
    etiqueta_estilo = ParagraphStyle('etq', parent=ESTILO_CELDA, fontSize=8, leading=10,
                                     textColor=colors.HexColor('#4B5E6D'))
    celdas = [[Paragraph(str(valor), cifra), Paragraph(etiqueta, etiqueta_estilo)] for etiqueta, valor in pares]
    t = Table([celdas], colWidths=[ANCHO / len(pares)] * len(pares))
    t.setStyle(TableStyle([('BOX', (0, 0), (-1, -1), 0.5, BORDE), ('INNERGRID', (0, 0), (-1, -1), 0.5, BORDE),
                           ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
    return t


def _quien(f):
    partes = [f.profesional.nombre_con_titulo if f.profesional else 'Todo el consultorio']
    if f.especialidad:
        partes.append(f.especialidad.nombre)
    return ' · '.join(partes)


def pdf(tenant, d):
    f = d['filtros']
    buffer = BytesIO()
    plantilla = documento(buffer, titulo_pdf='Reporte')
    h = encabezado_tenant(tenant, titulo='REPORTE DE GESTIÓN', subtitulo=f'{f.periodo_texto}<br/>{_quien(f)}',
                          profesional=f.profesional if not f.todo_el_consultorio or f.profesional else None)
    a = d['agenda']
    h += [_titulo('Actividad'), _kpis([
        ('Consultas', d['consultas']), ('Pacientes atendidos', d['pacientes_atendidos']),
        ('Pacientes nuevos', d['pacientes_nuevos']), ('Citas atendidas', a['atendidas']),
        ('Asistencia', f"{a['tasa_asistencia']}%" if a['tasa_asistencia'] is not None else '—')])]
    fi = d['finanzas']
    h += [_titulo('Ingresos (USD)'), _kpis([
        ('Facturado', usd(fi['facturado'])), ('Cobrado', usd(fi['cobrado'])), ('Costo', usd(fi['costo'])),
        ('Ingreso real', usd(fi['ingreso_real'])), ('Por cobrar', usd(fi['por_cobrar']))])]
    if d['mensual']:
        h += [_titulo('Por mes'), _tabla(['Mes', 'Consultas', 'Citas atendidas', 'Cobrado'],
                                         [[m['mes'], m['consultas'], m['citas_atendidas'], usd(m['cobrado'])]
                                          for m in d['mensual']], [60 * mm, 40 * mm, 43 * mm, 43 * mm], (1, 2, 3))]
    bloques = []
    if d['morbilidad']:
        bloques.append([_titulo('Diagnósticos más frecuentes (CIE-10)'),
                        _tabla(['Diagnóstico', 'Casos'], [[x['etiqueta'], x['valor']] for x in d['morbilidad']],
                               [156 * mm, 30 * mm], (1,))])
    bloques.append([_titulo('Agenda'), _tabla(
        ['Citas', 'Atendidas', 'No asistió', 'Canceladas', 'Pendientes', 'Sin cita previa'],
        [[a['total'], a['atendidas'], a['no_asistio'], a['canceladas'], a['pendientes'], a['sin_cita']]],
        [31 * mm] * 6, range(6))])
    if d['tipos']:
        bloques.append([_titulo('Tipo de consulta'), _tabla(['Tipo', 'Consultas'],
                                                            [[x['etiqueta'], x['valor']] for x in d['tipos']],
                                                            [156 * mm, 30 * mm], (1,))])
    if d['edades']:
        bloques.append([_titulo('Pacientes atendidos por edad y sexo'), _tabla(
            ['Grupo', 'Pacientes'], [[x['etiqueta'], x['valor']] for x in d['edades'] if x['valor']]
            + [[x['etiqueta'], x['valor']] for x in d['sexos']], [156 * mm, 30 * mm], (1,))])
    if fi['por_metodo']:
        bloques.append([_titulo('Cobrado por método'), _tabla(['Método', 'USD'],
                                                              [[x['etiqueta'], usd(x['valor'])] for x in fi['por_metodo']],
                                                              [146 * mm, 40 * mm], (1,))])
    if d['por_profesional']:
        bloques.append([_titulo('Por profesional'), _tabla(
            ['Profesional', 'Consultas', 'Pacientes', 'Citas atendidas', 'Cobrado'],
            [[x['profesional'].nombre_con_titulo, x['consultas'], x['pacientes'], x['citas_atendidas'],
              usd(x['cobrado'])] for x in d['por_profesional']],
            [66 * mm, 28 * mm, 28 * mm, 32 * mm, 32 * mm], (1, 2, 3, 4))])
    for b in bloques:
        h.append(KeepTogether(b))
    h += [Spacer(1, 4 * mm), Paragraph('<font size="7" color="#4B5E6D">Cobrado = pagos aplicados a cargos del período. '
                                       'Consultas anuladas excluidas. Pacientes nuevos = primera consulta en el período.'
                                       '</font>', ESTILOS['Normal'])]
    pie = pie_tenant(tenant)
    plantilla.build(h, onFirstPage=pie, onLaterPages=pie)
    return buffer.getvalue()


def excel(tenant, d, consultas):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    f = d['filtros']
    libro = Workbook()
    negrita = Font(bold=True, color='0C385B')
    relleno = PatternFill('solid', fgColor='E3EEF7')

    def hoja(nombre, encabezado, filas, primera=False):
        h = libro.active if primera else libro.create_sheet(nombre)
        h.title = nombre
        h.append([tenant.configuracion.nombre_display])
        h.append([f'Período: {f.periodo_texto} · {_quien(f)}'])
        h['A1'].font = Font(bold=True, size=13, color='145F8F')
        h.append([])
        h.append(encabezado)
        for celda in h[4]:
            celda.font, celda.fill = negrita, relleno
        for fila in filas:
            h.append(list(fila))
        for i, _ in enumerate(encabezado, 1):
            largo = max([len(str(encabezado[i - 1]))] + [len(str(fila[i - 1])) for fila in filas if len(fila) >= i])
            h.column_dimensions[get_column_letter(i)].width = min(max(12, largo + 2), 60)
        h.freeze_panes = 'A5'
        return h

    a, fi = d['agenda'], d['finanzas']
    hoja('Resumen', ['Indicador', 'Valor'], [
        ('Consultas', d['consultas']), ('Pacientes atendidos', d['pacientes_atendidos']),
        ('Pacientes nuevos', d['pacientes_nuevos']), ('Controles', d['controles']),
        ('Citas agendadas', a['total']), ('Citas atendidas', a['atendidas']), ('No asistió', a['no_asistio']),
        ('Canceladas', a['canceladas']), ('Atendidos sin cita previa', a['sin_cita']),
        ('Tasa de asistencia (%)', a['tasa_asistencia'] if a['tasa_asistencia'] is not None else ''),
        ('Facturado (USD)', float(fi['facturado'])), ('Cobrado (USD)', float(fi['cobrado'])),
        ('Costo (USD)', float(fi['costo'])), ('Ingreso real (USD)', float(fi['ingreso_real'])),
        ('Por cobrar (USD)', float(fi['por_cobrar'])), ('Cargos exonerados', fi['exonerados'])], primera=True)
    hoja('Por mes', ['Mes', 'Consultas', 'Citas atendidas', 'Cobrado (USD)'],
         [(m['mes'], m['consultas'], m['citas_atendidas'], float(m['cobrado'])) for m in d['mensual']])
    hoja('Morbilidad', ['Diagnóstico (CIE-10)', 'Casos'], [(x['etiqueta'], x['valor']) for x in d['morbilidad']])
    hoja('Pacientes', ['Grupo', 'Pacientes'], [(x['etiqueta'], x['valor']) for x in d['edades']]
         + [(x['etiqueta'], x['valor']) for x in d['sexos']])
    hoja('Cobros por método', ['Método', 'USD'], [(x['etiqueta'], float(x['valor'])) for x in fi['por_metodo']])
    if d['por_profesional']:
        hoja('Por profesional', ['Profesional', 'Especialidad', 'Consultas', 'Pacientes', 'Citas atendidas',
                                 'Cobrado (USD)'],
             [(x['profesional'].nombre_con_titulo, x['profesional'].especialidad_texto, x['consultas'], x['pacientes'],
               x['citas_atendidas'], float(x['cobrado'])) for x in d['por_profesional']])
    from django.utils import timezone
    detalle = hoja('Consultas', ['Fecha', 'Paciente', 'Documento', 'Profesional', 'Especialidad', 'Tipo', 'Estado',
                                 'Motivo', 'Diagnóstico'],
                   [(timezone.localtime(c.fecha_hora).strftime('%d/%m/%Y %I:%M %p'), c.paciente.nombre_completo,
                     c.paciente.documento, c.profesional.nombre_con_titulo,
                     c.especialidad.nombre if c.especialidad_id else '', c.get_tipo_display(), c.get_estado_display(),
                     c.motivo, c.diagnostico) for c in consultas])
    for fila in detalle.iter_rows(min_row=5):
        for celda in fila:
            celda.alignment = Alignment(vertical='top', wrap_text=celda.column > 7)
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
